import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from app.core.excecoes import ErroArquivoInvalido, ErroDocumentoSemTexto, ErroTipoNaoSuportado


@dataclass(frozen=True)
class TextoExtraido:
    """Texto de um arquivo, separado por página quando o formato tem páginas."""

    paginas: list[str]
    paginado: bool


def identificar_tipo(nome_arquivo: str) -> str:
    tipo = PurePath(nome_arquivo).suffix.lower().lstrip(".")
    if tipo not in _EXTRATORES:
        raise ErroTipoNaoSuportado("Formato não suportado. Envie um arquivo .txt, .pdf ou .docx.")
    return tipo


def extrair_texto(conteudo: bytes, tipo: str) -> TextoExtraido:
    extraido = _EXTRATORES[tipo](conteudo)
    paginas = [_normalizar(pagina) for pagina in extraido.paginas]
    if not any(paginas):
        raise ErroDocumentoSemTexto(
            "Não foi possível extrair texto do arquivo. Se for um PDF digitalizado "
            "(imagem), ele precisa passar por OCR antes do envio."
        )
    return TextoExtraido(paginas=paginas, paginado=extraido.paginado)


def _extrair_txt(conteudo: bytes) -> TextoExtraido:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Arquivos salvos pelo Bloco de Notas antigo no Windows costumam vir em cp1252.
        texto = conteudo.decode("cp1252", errors="replace")
    return TextoExtraido(paginas=[texto], paginado=False)


def _extrair_pdf(conteudo: bytes) -> TextoExtraido:
    try:
        leitor = PdfReader(io.BytesIO(conteudo))
        if leitor.is_encrypted:
            raise ErroArquivoInvalido("O PDF está protegido por senha.")
        paginas = [pagina.extract_text() or "" for pagina in leitor.pages]
    except PyPdfError as erro:
        raise ErroArquivoInvalido("O arquivo não é um PDF válido.") from erro
    return TextoExtraido(paginas=paginas, paginado=True)


def _extrair_docx(conteudo: bytes) -> TextoExtraido:
    try:
        documento = Document(io.BytesIO(conteudo))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as erro:
        raise ErroArquivoInvalido("O arquivo não é um DOCX válido.") from erro

    blocos = [paragrafo.text for paragrafo in documento.paragraphs]
    for tabela in documento.tables:
        for linha in tabela.rows:
            blocos.append(" | ".join(celula.text.strip() for celula in linha.cells))
    return TextoExtraido(paginas=["\n\n".join(blocos)], paginado=False)


def _normalizar(texto: str) -> str:
    texto = texto.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    texto = re.sub(r"[ \t\xa0]+", " ", texto)
    texto = re.sub(r" ?\n ?", "\n", texto)
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    return texto.strip()


_EXTRATORES = {"txt": _extrair_txt, "pdf": _extrair_pdf, "docx": _extrair_docx}
