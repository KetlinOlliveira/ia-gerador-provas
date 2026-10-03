import io
import re
import unicodedata
from dataclasses import dataclass
from enum import StrEnum

from docx import Document
from docx.shared import Pt
from fpdf import FPDF

from app.esquemas.provas import ProvaDetalhe
from app.exportacao.conteudo import Bloco, Versao, montar_blocos


class Formato(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    MARKDOWN = "md"


@dataclass(frozen=True)
class ArquivoExportado:
    conteudo: bytes
    tipo_midia: str
    nome: str


def exportar(prova: ProvaDetalhe, formato: Formato, versao: Versao) -> ArquivoExportado:
    blocos = montar_blocos(prova, versao)
    sufixo = "-gabarito" if versao == Versao.PROFESSOR else ""
    nome = f"{_nome_de_arquivo(prova.titulo)}{sufixo}.{formato.value}"

    if formato == Formato.PDF:
        return ArquivoExportado(_para_pdf(blocos), "application/pdf", nome)
    if formato == Formato.DOCX:
        tipo = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        return ArquivoExportado(_para_docx(blocos), tipo, nome)
    return ArquivoExportado(_para_markdown(blocos).encode(), "text/markdown; charset=utf-8", nome)


def _nome_de_arquivo(titulo: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", titulo).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-zA-Z0-9]+", "-", sem_acento).strip("-").lower() or "prova"


# ---------- Markdown ----------


def _para_markdown(blocos: list[Bloco]) -> str:
    linhas = []
    for bloco in blocos:
        match bloco.tipo:
            case "titulo":
                linhas.append(f"# {bloco.texto}\n")
            case "subtitulo":
                linhas.append(f"_{bloco.texto}_\n")
            case "questao":
                linhas.append(f"\n## {bloco.texto}\n")
            case "item":
                linhas.append(f"- {bloco.texto}")
            case "gabarito":
                linhas.append(f"\n**{bloco.texto}**\n")
            case "linhas":
                linhas.append("\n" + "\n".join(["_" * 60] * 5) + "\n")
            case _:
                linhas.append(f"{bloco.texto}\n")
    return "\n".join(linhas).strip() + "\n"


# ---------- DOCX ----------


def _para_docx(blocos: list[Bloco]) -> bytes:
    documento = Document()
    documento.styles["Normal"].font.size = Pt(11)
    for bloco in blocos:
        match bloco.tipo:
            case "titulo":
                documento.add_heading(bloco.texto, level=0)
            case "subtitulo":
                documento.add_paragraph(bloco.texto).runs[0].italic = True
            case "questao":
                documento.add_heading(bloco.texto, level=2)
            case "item":
                documento.add_paragraph(bloco.texto, style="List Bullet")
            case "gabarito":
                documento.add_paragraph().add_run(bloco.texto).bold = True
            case "linhas":
                for _ in range(5):
                    documento.add_paragraph("_" * 80)
            case _:
                documento.add_paragraph(bloco.texto)
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


# ---------- PDF ----------

# As fontes embutidas do PDF só cobrem Latin-1. Acentos do português estão lá; aspas
# curvas, travessões e hifens especiais que os modelos gostam de usar, não.
_TROCAS_PDF = str.maketrans(
    {
        "‐": "-",
        "‑": "-",
        "‒": "-",
        "–": "-",
        "—": "-",
        "‘": "'",
        "’": "'",
        "“": '"',
        "”": '"',
        "…": "...",
        "•": "-",
        "·": "-",
        "→": "->",
    }
)


def _texto_pdf(texto: str) -> str:
    return texto.translate(_TROCAS_PDF).encode("latin-1", "replace").decode("latin-1")


def _para_pdf(blocos: list[Bloco]) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(18, 18, 18)
    pdf.add_page()

    def escrever(texto: str, estilo: str = "", tamanho: float = 11, altura: float = 6) -> None:
        pdf.set_font("Helvetica", estilo, tamanho)
        pdf.multi_cell(0, altura, _texto_pdf(texto), new_x="LMARGIN", new_y="NEXT")

    for bloco in blocos:
        match bloco.tipo:
            case "titulo":
                escrever(bloco.texto, "B", 18, 9)
            case "subtitulo":
                escrever(bloco.texto, "I", 11)
                pdf.ln(3)
            case "questao":
                pdf.ln(4)
                escrever(bloco.texto, "B", 12, 7)
            case "item":
                pdf.set_x(pdf.l_margin + 5)
                pdf.set_font("Helvetica", "", 11)
                pdf.multi_cell(0, 6, _texto_pdf(bloco.texto), new_x="LMARGIN", new_y="NEXT")
            case "gabarito":
                pdf.ln(1)
                escrever(bloco.texto, "B")
            case "linhas":
                for _ in range(5):
                    pdf.ln(7)
                    y = pdf.get_y()
                    pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(2)
            case _:
                escrever(bloco.texto)
    return bytes(pdf.output())
