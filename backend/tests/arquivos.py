"""Gera arquivos de exemplo em memória, para os testes não dependerem de binários no repositório."""

import io

from docx import Document


def criar_docx(paragrafos: list[str], tabela: list[list[str]] | None = None) -> bytes:
    documento = Document()
    for paragrafo in paragrafos:
        documento.add_paragraph(paragrafo)
    if tabela:
        grade = documento.add_table(rows=len(tabela), cols=len(tabela[0]))
        for linha, valores in zip(grade.rows, tabela, strict=True):
            for celula, valor in zip(linha.cells, valores, strict=True):
                celula.text = valor
    saida = io.BytesIO()
    documento.save(saida)
    return saida.getvalue()


def criar_pdf(paginas: list[str]) -> bytes:
    """Monta um PDF mínimo com uma linha de texto (ASCII) por página."""
    total = len(paginas)
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count %d /Kids [%s] >>"
        % (total, b" ".join(b"%d 0 R" % (4 + 2 * i) for i in range(total))),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for indice, texto in enumerate(paginas):
        conteudo = b"BT /F1 12 Tf 72 720 Td (%s) Tj ET" % texto.encode("ascii")
        objetos.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * indice)
        )
        objetos.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(conteudo), conteudo))

    pdf = b"%PDF-1.4\n"
    posicoes = []
    for numero, objeto in enumerate(objetos, start=1):
        posicoes.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (numero, objeto)

    inicio_xref = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % posicao for posicao in posicoes)
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objetos) + 1,
        inicio_xref,
    )
    return pdf
